# -*- coding: utf-8 -*-
import sys
import shutil
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

p1 = r"C:\Users\dienv\Desktop\Data4Life\02_THUYET_TRINH_10_PHUT\SLIDE_THUYET_TRINH_10_PHUT_SOS_VIET_NAM_2026.pptx"
p2 = r"C:\Users\dienv\Desktop\docs\thuyết trình\SLIDE_THUYET_TRINH_10_PHUT_SOS_VIET_NAM_2026.pptx"

prs = Presentation(p1)
slide11 = prs.slides[10]
shape3 = slide11.shapes[3]

tf = shape3.text_frame
tf.clear()

p0 = tf.paragraphs[0]
p0.text = "“Trong kỷ nguyên chuyển đổi số quốc gia, khi dữ liệu số có thể đi nhanh hơn cả tiếng còi xe cứu thương —"
p0.font.name = "Arial"
p0.font.size = Pt(15)
p0.font.bold = True
p0.font.italic = True
p0.font.color.rgb = RGBColor(255, 230, 100)

p1_para = tf.add_paragraph()
p1_para.text = "Liệu chúng ta có thể cho phép bất kỳ người dân nào phải đơn độc trong 5 phút sinh tử của cuộc đời mình chỉ vì một rào cản kết nối?”"
p1_para.font.name = "Arial"
p1_para.font.size = Pt(15)
p1_para.font.bold = True
p1_para.font.italic = True
p1_para.font.color.rgb = RGBColor(255, 255, 255)

prs.save(p1)
shutil.copy2(p1, p2)
print("[OK] Updated Slide 11 successfully.")
