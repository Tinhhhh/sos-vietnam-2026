import os
import sys
import shutil
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.stdout.reconfigure(encoding='utf-8')

DOCX_PATH = r"C:\Users\dienv\Desktop\docs\thuyết trình\DE_XUAT_GIAI_PHAP_SOS_VIET_NAM_2026_DATA_FOR_LIFE_CAP_NHAT.docx"
BACKUP_PATH = r"C:\Users\dienv\Desktop\docs\thuyết trình\DE_XUAT_GIAI_PHAP_SOS_VIET_NAM_2026_DATA_FOR_LIFE_CAP_NHAT.docx.bak"

# 1. Create backup
if not os.path.exists(BACKUP_PATH):
    shutil.copy2(DOCX_PATH, BACKUP_PATH)
    print(f"[OK] Created backup at: {BACKUP_PATH}")

doc = docx.Document(DOCX_PATH)

# Find placeholder paragraph 17: "(*Bổ sung cách hệ thống..."
target_idx = None
for i, p in enumerate(doc.paragraphs):
    if "(*Bổ sung" in p.text:
        target_idx = i
        break

if target_idx is None:
    print("[ERROR] Could not find placeholder paragraph in docx!")
    exit(1)

print(f"[OK] Found target placeholder at index {target_idx}")
target_p = doc.paragraphs[target_idx]

# Helper to format run
def format_run(run, bold=False, italic=False, size=13, font_name="Times New Roman", color=None):
    run.font.name = font_name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color

# Content blocks to insert before target_p
section_2_content = [
    {
        "type": "subheading",
        "title": "2.1. Tóm Tắt Giải Pháp — 3 Năng Lực Trọng Tâm Dự Án Đang Làm Được:"
    },
    {
        "type": "item",
        "bold_prefix": "1. Hợp nhất dữ liệu không gian C4ISR và điều phối liên ngành thời gian thực trên bản đồ GIS 3D (MapLibre GIS 3D): ",
        "text": "Xóa bỏ triệt để tình trạng cát cứ, phân mảnh thông tin giữa các tổng đài riêng lẻ (113, 114, 115 và cứu hộ giao thông). Hệ thống ứng dụng công nghệ định vị W3C Geolocation kết hợp thuật toán lọc Kalman Filter (đạt sai số thực nghiệm < 5m), kết nối đồng bộ trực tiếp lưới 34 tỉnh/thành phố, 3.321 xã/phường từ cổng bando.com.vn và cơ sở dữ liệu 192 trạm tác chiến trên toàn quốc (đồng bộ hai chiều Excel 4 sheet). Khi phát sinh sự cố, máy chủ tự động khóa camera bản đồ 3D tại hiện trường (FlyTo), phân luồng chính xác theo địa giới hành chính và tối ưu hóa tuyến đường ngắn nhất cho lực lượng cơ động tiếp cận hiện trường trong 'khung giờ vàng'."
    },
    {
        "type": "item",
        "bold_prefix": "2. Cổng cứu nạn PWA 1-Chạm siêu nhẹ (< 1.8 KB) cho người yếu thế & Cơ chế chuyển cấp tự động (Timeout Escalation): ",
        "text": "Giải quyết dứt điểm điểm nghẽn đề tài ASXH-07: Người cao tuổi, người khuyết tật, người trong tình huống hoảng loạn không cần cài đặt phần mềm từ kho ứng dụng hay đăng nhập tài khoản phức tạp. Người dân chỉ cần mở web PWA 1-chạm dưới 1 giây, hoạt động bền bỉ cả khi mất sóng Internet nhờ công nghệ Service Worker và lưu trữ ngoại tuyến IndexedDB; hỗ trợ gọi thoại/video WebRTC trực tiếp và truyền tức thì viễn trắc thiết bị. Đặc biệt, hệ thống tích hợp bộ đo thời gian tự động: Nếu sự vụ nguy cấp quá 15 phút mà cấp cơ sở (xã/phường) chưa xử lý xong, hệ thống tự động leo thang (Timeout Escalation) chuyển quyền chỉ huy lên cấp Tỉnh/Thành phố và Trung tâm Tác chiến Quốc gia để điều động chi viện, kiên quyết không bỏ sót bất kỳ sinh mạng nào."
    },
    {
        "type": "item",
        "bold_prefix": "3. Lá chắn an ninh mạng Layer 7, trích xuất chứng cứ số OSINT và chế tài xử lý tin báo khống theo Nghị định 144/2021/NĐ-CP: ",
        "text": "Bảo vệ hệ thống bằng tường lửa WAF Layer 7 chặn 27 chữ ký AI Scraper Bot cào dữ liệu, chống giả lập GPS (GPS Spoofing), bẫy Honeypot và kiểm soát phân quyền Zero-Trust RBAC. Khi phát hiện hiện trường giả hoặc hành vi quấy rối, cán bộ trực ban chỉ cần một cú nhấp chuột để kích hoạt mô-đun Cyber Forensics: Hệ thống tự động trích xuất địa chỉ IP/ISP, User-Agent, dấu vân tay thiết bị, bộ lệnh Linux điều tra, tự động điền vào Biên bản vi phạm hành chính chuẩn mẫu Nghị định số 144/2021/NĐ-CP và đưa thiết bị vào Danh sách đen (Blacklist) toàn quốc, triệt tiêu lãng phí hàng chục tỷ đồng ngân sách cho các đợt xuất kích xe công vụ vô ích."
    },
    {
        "type": "subheading",
        "title": "2.2. Tầm Nhìn Chiến Lược — 2 Trục Mở Rộng Quy Mô (Scope) Tương Lai:"
    },
    {
        "type": "item",
        "bold_prefix": "1. Tích hợp sâu vào Hệ sinh thái Định danh Quốc gia VNeID & CSDL Hồ sơ Sức khỏe Điện tử (Đề án 06/CP): ",
        "text": "Mở rộng liên thông API trực tiếp với Cơ sở dữ liệu quốc gia về dân cư qua tài khoản VNeID mức độ 2. Khi công dân phát tín hiệu SOS, hệ thống tự động đối soát thông tin nhân thân, nhóm máu, tiền sử bệnh án/dị ứng thuốc, tình trạng khuyết tật và số liên lạc khẩn cấp của người thân từ CSDL Y tế Quốc gia. Bác sĩ trực ban 115 có thể chỉ định phác đồ cấp cứu chuẩn xác ngay trên xe cứu thương trong 'khung giờ vàng', hiện thực hóa mục tiêu lấy dữ liệu số phụng sự trực tiếp mạng sống con người theo tinh thần Đề án 06/CP."
    },
    {
        "type": "item",
        "bold_prefix": "2. Mở rộng cảm biến IoT Cảnh báo sớm thiên tai, Camera AI giao thông và Mạng lưới cứu trợ ASEAN ('Build Together'): ",
        "text": "Mở rộng tích hợp mạng lưới cảm biến IoT đo mực nước ngập đô thị/sạt lở và cảm biến khói báo cháy tự động tại rừng và khu công nghiệp nhằm tự động phát lệnh cứu hộ máy-sang-máy (Autonomous M2M Emergency Dispatch) mà không cần con người thao tác; đồng thời kết nối luồng camera AI giao thông để tự động phát hiện tai nạn giao thông nghiêm trọng. Thiết lập trục điều phối tác chiến liên tỉnh toàn vùng Đồng bằng sông Cửu Long và kết nối hạ tầng dữ liệu cứu trợ khẩn cấp khu vực các nước ASEAN theo đúng sứ mệnh 'Build Together' của Cuộc thi Quốc tế Data for Life."
    }
]

# Insert paragraphs before target_p
for item in section_2_content:
    new_p = target_p.insert_paragraph_before()
    new_p.paragraph_format.line_spacing = 1.2
    new_p.paragraph_format.space_after = Pt(6)
    new_p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    if item["type"] == "subheading":
        run = new_p.add_run(item["title"])
        format_run(run, bold=True, size=13, color=RGBColor(0, 51, 102))  # Dark Blue
        new_p.paragraph_format.space_before = Pt(8)
    elif item["type"] == "item":
        run_bold = new_p.add_run(item["bold_prefix"])
        format_run(run_bold, bold=True, size=12)
        run_text = new_p.add_run(item["text"])
        format_run(run_text, bold=False, size=12)

# Remove the placeholder paragraph
target_p._element.getparent().remove(target_p._element)
print("[OK] Removed placeholder paragraph.")

# Remove leftover draft duplicate headings (paragraphs between section 2 and section 3)
# We inspect paragraphs starting with '3. Tính năng' or '3.1 Vai trò'
to_remove = []
for p in doc.paragraphs:
    txt = p.text.strip()
    if txt in ["3. Tính năng của hệ thống.", "3.1 Vai trò.", "3.1 Vai trò.\n\n3.1 Vai trò."]:
        to_remove.append(p)

for p in to_remove:
    p._element.getparent().remove(p._element)
    print(f"[OK] Cleaned draft duplicate: '{p.text[:30]}...'")

# Save document to completed path first (avoids lock by Word)
COMPLETED_PATH = r"C:\Users\dienv\Desktop\docs\thuyết trình\DE_XUAT_GIAI_PHAP_SOS_VIET_NAM_2026_DATA_FOR_LIFE_CAP_NHAT_HOAN_THIEN.docx"
doc.save(COMPLETED_PATH)
print(f"[SUCCESS] Saved completed docx file to: {COMPLETED_PATH}")

# Also try to overwrite original DOCX_PATH if not locked
try:
    doc.save(DOCX_PATH)
    print(f"[SUCCESS] Overwrote original file: {DOCX_PATH}")
except PermissionError:
    print(f"[NOTICE] Original file is currently open in Microsoft Word. Saved to '{COMPLETED_PATH}' so you can review without losing Word state.")
