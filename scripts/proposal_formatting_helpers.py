import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

sys.stdout.reconfigure(encoding='utf-8')

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith('shd'):
            tcPr.remove(child)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    for child in list(tblPr):
        if child.tag.endswith('tblBorders'):
            tblPr.remove(child)
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def remove_table_borders(table):
    tblPr = table._tbl.tblPr
    for child in list(tblPr):
        if child.tag.endswith('tblBorders'):
            tblPr.remove(child)
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="none" w:sz="0" w:space="0" w:color="auto"/>\n'
        f'  <w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>\n'
        f'  <w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/>\n'
        f'  <w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>\n'
        f'  <w:insideH w:val="none" w:sz="0" w:space="0" w:color="auto"/>\n'
        f'  <w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def set_cell_margins(cell, top=70, bottom=70, left=110, right=110):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>\n'
        f'  <w:top w:w="{top}" w:type="dxa"/>\n'
        f'  <w:bottom w:w="{bottom}" w:type="dxa"/>\n'
        f'  <w:left w:w="{left}" w:type="dxa"/>\n'
        f'  <w:right w:w="{right}" w:type="dxa"/>\n'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def prevent_row_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

def format_p(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=2.5, line_spacing=1.15, first_line_indent=0.0, keep_with_next=False):
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if keep_with_next:
        p.paragraph_format.keep_with_next = True
    if first_line_indent > 0:
        p.paragraph_format.first_line_indent = Cm(first_line_indent)

def add_p(doc, text="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=2.5, line_spacing=1.15, first_line_indent=0.75, font_size=12, italic=False, bold=False):
    p = doc.add_paragraph()
    format_p(p, align, space_before, space_after, line_spacing, first_line_indent)
    if text:
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(font_size)
        r.italic = italic
        r.bold = bold
        r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_p_runs(doc, runs_data, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=2.5, line_spacing=1.15, first_line_indent=0.75):
    p = doc.add_paragraph()
    format_p(p, align, space_before, space_after, line_spacing, first_line_indent)
    for text, bold, italic, size in runs_data:
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(size if size else 12)
        r.bold = bold
        r.italic = italic
        r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_h1(doc, text):
    p = doc.add_paragraph()
    format_p(p, WD_ALIGN_PARAGRAPH.LEFT, space_before=5, space_after=2, line_spacing=1.15, first_line_indent=0, keep_with_next=True)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12.5)
    r.bold = True
    r.font.color.rgb = RGBColor(0, 32, 96) # Dark Navy formal
    return p

def add_h2(doc, text):
    p = doc.add_paragraph()
    format_p(p, WD_ALIGN_PARAGRAPH.LEFT, space_before=3.5, space_after=1.5, line_spacing=1.15, first_line_indent=0, keep_with_next=True)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.bold = True
    r.font.color.rgb = RGBColor(15, 23, 42)
    return p

def add_h3(doc, text):
    p = doc.add_paragraph()
    format_p(p, WD_ALIGN_PARAGRAPH.LEFT, space_before=2.5, space_after=1, line_spacing=1.15, first_line_indent=0, keep_with_next=True)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.bold = True
    r.italic = True
    r.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_image_placeholder(doc, fig_num, fig_title, fig_desc=""):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    prevent_row_split(table.rows[0])
    
    cell = table.cell(0, 0)
    cell.width = Cm(16.5)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=80, bottom=80, left=140, right=140)
    
    # Border: subtle single border
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="94A3B8"/>\n'
        f'  <w:left w:val="single" w:sz="6" w:space="0" w:color="94A3B8"/>\n'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="94A3B8"/>\n'
        f'  <w:right w:val="single" w:sz="6" w:space="0" w:color="94A3B8"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    r1 = p.add_run("[VỊ TRÍ BỔ SUNG HÌNH ẢNH MINH HỌA - TÀI LIỆU CHÍNH THỨC]")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(10)
    r1.bold = True
    r1.font.color.rgb = RGBColor(71, 85, 105)
    
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(2)
    r2 = p2.add_run(f"Hình {fig_num}: {fig_title}")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(10.5)
    r2.bold = True
    r2.italic = True
    r2.font.color.rgb = RGBColor(15, 23, 42)
    
    if fig_desc:
        p3 = cell.add_paragraph()
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p3.paragraph_format.space_before = Pt(0)
        p3.paragraph_format.space_after = Pt(3)
        r3 = p3.add_run(f"({fig_desc})")
        r3.font.name = "Times New Roman"
        r3.font.size = Pt(9.5)
        r3.italic = True
        r3.font.color.rgb = RGBColor(100, 116, 139)
        
    sp = doc.add_paragraph()
    format_p(sp, space_before=0, space_after=2.5, line_spacing=1.0)
    return table

def style_cell_run(cell, text, bold=False, italic=False, size=10.5, align=WD_ALIGN_PARAGRAPH.LEFT, color=RGBColor(0,0,0)):
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    r.font.color.rgb = color
    return r

print("Loaded enhanced proposal formatting helpers successfully!")
