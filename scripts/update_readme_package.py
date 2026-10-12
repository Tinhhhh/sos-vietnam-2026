import os
import sys
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import win32com.client
import pythoncom

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_readme_docx():
    doc = docx.Document()
    
    # Page setup
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(0.8)
        
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    normal_style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(12)
    r1 = p_title.add_run("TỔNG QUAN BỘ HỒ SƠ DỰ THI DATA FOR LIFE 2026\n(TOP 60 VÒNG CHINH PHỤC)")
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(15)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

    # General info
    infos = [
        ("Dự án: ", "SOS VIỆT NAM 2026 – Nền tảng dữ liệu và điều phối cứu nạn khẩn cấp đa lực lượng"),
        ("Bài toán: ", "ASXH-07 (Ứng cứu khẩn cấp, cứu nạn cứu hộ và an sinh xã hội)"),
        ("Nhóm tác giả: ", "03 thành viên chính thức:"),
    ]
    for label, val in infos:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        r_lbl = p.add_run(label)
        r_lbl.font.bold = True
        r_lbl.font.name = 'Times New Roman'
        r_val = p.add_run(val)
        r_val.font.name = 'Times New Roman'

    members = [
        ("Điền Trần Vĩnh An", "(Trưởng nhóm / Chủ nhiệm đề tài / Đội trưởng - Đóng góp: 60%)", "Chỉ đạo nghiên cứu tổng thể; ứng dụng trí tuệ nhân tạo (AI-assisted / Vibe Coding) tối ưu phát triển phần mềm; thiết kế kiến trúc hệ thống C4ISR; xây dựng động cơ bản đồ GIS 3D; thuật toán lọc nhiễu Kalman Filter và số hóa Nghị định 30/2020/NĐ-CP."),
        ("[Họ và tên - Tạm chưa điền / Bổ sung sau]", "(Thành viên - Đóng góp: 25%)", "Kỹ thuật lập trình, rà soát mã nguồn, kiểm thử phát hiện và khắc phục sự cố, fix code bug, tối ưu Web PWA và CSDL."),
        ("[Họ và tên - Tạm chưa điền / Bổ sung sau]", "(Thành viên - Đóng góp: 15%)", "Biên tập nội dung, đồ họa UI/UX, sản xuất và hậu kỳ video demo, thiết kế slide thuyết trình và hồ sơ truyền thông."),
    ]
    for idx, (name, role, desc) in enumerate(members, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        r_num = p.add_run(f"{idx}. ")
        r_num.font.bold = True
        r_num.font.name = 'Times New Roman'
        r_name = p.add_run(f"{name} ")
        r_name.font.bold = True
        r_name.font.name = 'Times New Roman'
        r_role = p.add_run(f"{role} – ")
        r_role.font.italic = True
        r_role.font.name = 'Times New Roman'
        r_desc = p.add_run(desc)
        r_desc.font.name = 'Times New Roman'

    extra_infos = [
        ("Thời gian cập nhật: ", "Tháng 10 năm 2026 (Số liệu địa giới hành chính cập nhật mới nhất từ bando.com.vn và danh bạ 874 trạm tác chiến trên toàn quốc)"),
        ("Cơ quan chỉ đạo & Tổ chức: ", "Cục Cảnh sát QLHC về TTXH (C06 - Bộ Công an) | ĐH Bách Khoa Hà Nội | Đài Truyền hình Việt Nam (VTV)"),
    ]
    for label, val in extra_infos:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        r_lbl = p.add_run(label)
        r_lbl.font.bold = True
        r_lbl.font.name = 'Times New Roman'
        r_val = p.add_run(val)
        r_val.font.name = 'Times New Roman'

    # Divider
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(6)
    p_div.paragraph_format.space_after = Pt(6)
    r_div = p_div.add_run("─" * 55)
    r_div.font.color.rgb = RGBColor(0xaa, 0xaa, 0xaa)

    # Section 1: Timeline
    p_s1 = doc.add_paragraph()
    p_s1.paragraph_format.space_before = Pt(6)
    p_s1.paragraph_format.space_after = Pt(4)
    r_s1 = p_s1.add_run("📌 CÁC MỐC THỜI GIAN QUAN TRỌNG")
    r_s1.font.name = 'Times New Roman'
    r_s1.font.size = Pt(12)
    r_s1.font.bold = True
    r_s1.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

    timeline_items = [
        ("Kênh Discord chính thức: ", "https://discord.gg/E73gt5aqz (Bắt buộc theo dõi hàng ngày)."),
        ("23/10/2026 (Thứ 6): ", "Chương trình Mentoring (Sáng: Trực tiếp tại ĐH Bách Khoa Hà Nội; Chiều: Online)."),
        ("07/11/2026 (Thứ 7): ", "Đánh giá Vòng Chinh phục (10 phút Thuyết trình & Demo + 05 phút Q&A Phản biện) để chọn TOP 30."),
        ("Đầu tháng 12/2026: ", "Vòng Triển lãm Techfest tại Phố đi bộ Hồ Gươm, Hà Nội (Chấm điểm trực tiếp tại gian hàng) để chọn TOP 10."),
        ("Vòng Chung kết: ", "Tranh tài và ghi hình trực tiếp tại Đài Truyền hình Việt Nam (VTV)."),
    ]
    for lbl, val in timeline_items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        r_b = p.add_run("• ")
        r_b.font.name = 'Times New Roman'
        r_l = p.add_run(lbl)
        r_l.font.bold = True
        r_l.font.name = 'Times New Roman'
        r_v = p.add_run(val)
        r_v.font.name = 'Times New Roman'

    # Divider
    p_div2 = doc.add_paragraph()
    p_div2.paragraph_format.space_before = Pt(6)
    p_div2.paragraph_format.space_after = Pt(6)
    r_div2 = p_div2.add_run("─" * 55)
    r_div2.font.color.rgb = RGBColor(0xaa, 0xaa, 0xaa)

    # Section 2: Directory structure
    p_s2 = doc.add_paragraph()
    p_s2.paragraph_format.space_before = Pt(6)
    p_s2.paragraph_format.space_after = Pt(4)
    r_s2 = p_s2.add_run("📂 BẢN ĐỒ THƯ MỤC HỒ SƠ CHÍNH THỨC TẠI C:\\Users\\dienv\\Desktop\\Data4Life")
    r_s2.font.name = 'Times New Roman'
    r_s2.font.size = Pt(12)
    r_s2.font.bold = True
    r_s2.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

    dirs = [
        ("📁 00_THONG_BAO_VA_QUY_CHE", [
            ("THONG_BAO_TOP60_VA_HUONG_DAN_VONG_CHINH_PHUC_DATA_FOR_LIFE_2026.docx & .pdf", "Toàn văn thông báo chính thức, timeline, quy chế và các lưu ý bắt buộc."),
            ("[DATA FOR LIFE 2026] THÔNG BÁO KẾT QUẢ TOP 60...eml", "Email gốc từ Ban Tổ chức cuộc thi."),
        ]),
        ("📁 01_DE_XUAT_GIAI_PHAP", [
            ("DE_XUAT_GIAI_PHAP_SOS_VIET_NAM_2026_DATA_FOR_LIFE.docx & .pdf", "Bản Đề xuất giải pháp khoa học & công nghệ chính thức nộp Hội đồng Giám khảo Vòng Chinh Phục (Top 60 Toàn quốc). Độ dài đúng 07 trang (chuẩn tiêu chí dưới 10 trang), thể thức chuẩn Nghị định 30/2020/NĐ-CP (100% font Times New Roman, ký hiệu văn bản phân cấp động, loại bỏ hoàn toàn cấp CATT), bám sát trọn vẹn 8 trọng tâm đánh giá của Ban Tổ chức, 4 khung vị trí bổ sung ảnh minh họa tinh gọn, bảng phân định rõ ràng giữa tính năng hiện tại (TRL 7 trên Web, 80% mức độ hoàn thiện) và định hướng tương lai (R&D Roadmap), cùng danh mục 03 thành viên. Quản lý mạng lưới 874 trạm tác chiến trên 34 tỉnh thành và 3.321 xã phường sau sáp nhập."),
            ("SOS_VIETNAM_2026.docx & .pdf", "Toàn văn Báo cáo Thuyết minh Đề tài Nghiên cứu Khoa học và Phát triển Công nghệ toàn diện (Cơ sở khoa học và căn cứ gốc của hệ thống, cập nhật tháng 10/2026, nhóm 03 thành viên, bổ sung bảng thể thức văn bản chuẩn NĐ 30, VieNeu-TTS 4 giọng nữ 4 vùng miền tăng tốc 1.2x, WebRTC Video Call 2 chiều, Quả cầu SOS Orb căn giữa tâm sóng)."),
        ]),
        ("📁 02_THUYET_TRINH_10_PHUT", [
            ("KICH_BAN_THUYET_TRINH_10_PHUT_DE_HOC_THUOC.docx & .pdf", "Kịch bản video giới thiệu app & thuyết minh (chuẩn 10 phút), vào thẳng vấn đề '5 phút vàng' ở đầu video không chào hỏi khách sáo, trình bày trọn vẹn toàn bộ tính năng và kết thúc bằng câu hỏi đắt giá, sâu sắc về cứu nạn và chuyển đổi số."),
            ("BAI_THUYET_TRINH_8_TIEU_CHI_SOS_VIET_NAM_2026.docx & .pdf", "Toàn văn bài thuyết trình giải đáp chi tiết 8 tiêu chí chấm thi chính thức của Hội đồng Giám khảo."),
            ("BAI_DOC_THUYET_TRINH_SOS_VIET_NAM_2026.mp3", "File giọng đọc thuyết minh video chuẩn Studio (Nam Minh Neural), giọng trầm ấm, dõng dạc, đồng bộ 100% với kịch bản mới."),
            ("SLIDE_THUYET_TRINH_10_PHUT_SOS_VIET_NAM_2026.pptx & .pdf", "Slide trình chiếu 16:9 chuyên nghiệp (11 slide bám sát bài nói, phân định rõ tính năng hiện tại và tương lai, huy hiệu 80% TRL 7, số liệu 874 trạm tác chiến)."),
            ("SLIDE_THUYET_TRINH_8_TIEU_CHI_SOS_VIET_NAM_2026.pptx, .html & .pdf", "Bộ slide trình chiếu và nền tảng trình chiếu HTML chuyên biệt trả lời trực diện 8 tiêu chí đánh giá Data for Life 2026."),
        ]),
        ("📁 03_PHAN_BIEN_5_PHUT_BGK", [
            ("BO_CAU_HOI_TRA_LOI_5_PHUT_BAN_GIAM_KHAO.docx & .pdf", "9 câu hỏi cốt lõi kèm đáp án phản biện chuẩn xác (Hạ tầng viễn thông, OSRM 2 đường đi, CSDL địa giới T10/2026 với 874 trạm tác chiến, chống spam, bảo mật Nghị định 13, Stress test 10.000 req/s, TRL 7 80%, Đề án 06 và phân định tính năng hiện tại vs tương lai)."),
        ]),
        ("📁 04_KIEN_TRUC_VA_SO_DO", [
            ("role-hierarchy.flow.svg & .pdf", "Sơ đồ luồng phân cấp và phân quyền tác chiến 6 tầng."),
            ("role-hierarchy.flow-map.md & .flow-verification.md", "Báo cáo đối soát ngữ nghĩa và kiểm thử luồng."),
        ]),
        ("📁 05_TAI_LIEU_VAN_HANH_HUONG_DAN", [
            ("HUONG_DAN_SU_DUNG_SOS_VIET_NAM_2026_CAN_BO.docx & .pdf", "Cẩm nang vận hành dành cho Cán bộ trực ban chỉ huy (Cập nhật giao diện WebRTC Video Call, xuất biên bản chuẩn Nghị định 30/2020/NĐ-CP, phân cấp số văn bản động, không dùng CATT)."),
            ("HUONG_DAN_SU_DUNG_SOS_VIET_NAM_2026_NGUOI_DAN.docx & .pdf", "Cẩm nang sử dụng ứng dụng 1-chạm dành cho công dân (Cập nhật phím Master SOS Orb căn giữa tâm sóng video, Trợ lý giọng nói AI 4 vùng miền VieNeu-TTS tốc độ 1.2x)."),
            ("Tai_Lieu_Tong_Quan_Va_Huong_Dan_Van_Hanh_SOS_Viet_Nam_2026.docx & .pdf", "Tài liệu tổng quan kiến trúc kỹ thuật hệ thống."),
        ]),
    ]

    for dir_name, files in dirs:
        p_dir = doc.add_paragraph()
        p_dir.paragraph_format.space_before = Pt(6)
        p_dir.paragraph_format.space_after = Pt(2)
        r_d = p_dir.add_run(dir_name)
        r_d.font.bold = True
        r_d.font.name = 'Times New Roman'
        r_d.font.size = Pt(11.5)
        r_d.font.color.rgb = RGBColor(0x00, 0x44, 0x88)

        for fname, desc in files:
            p_f = doc.add_paragraph()
            p_f.paragraph_format.left_indent = Inches(0.25)
            p_f.paragraph_format.space_after = Pt(2)
            p_f.paragraph_format.line_spacing = 1.15
            r_dot = p_f.add_run("• ")
            r_dot.font.name = 'Times New Roman'
            r_fn = p_f.add_run(f"{fname}: ")
            r_fn.font.bold = True
            r_fn.font.name = 'Times New Roman'
            r_dc = p_f.add_run(desc)
            r_dc.font.name = 'Times New Roman'

    docx_path = r"C:\Users\dienv\Desktop\Data4Life\README_TONG_QUAN_BO_HO_SO_DATA4LIFE_2026.docx"
    doc.save(docx_path)
    print(f"Saved DOCX: {docx_path}")

    # Export to PDF via Word COM
    pdf_path = r"C:\Users\dienv\Desktop\Data4Life\README_TONG_QUAN_BO_HO_SO_DATA4LIFE_2026.pdf"
    try:
        pythoncom.CoInitialize()
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        wdoc = word.Documents.Open(os.path.abspath(docx_path))
        wdoc.SaveAs(os.path.abspath(pdf_path), FileFormat=17)
        wdoc.Close(False)
        word.Quit()
        print(f"Saved PDF: {pdf_path}")
    except Exception as e:
        print(f"PDF export error: {e}")

if __name__ == "__main__":
    create_readme_docx()
